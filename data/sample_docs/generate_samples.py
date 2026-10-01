"""
Sample Document Generator for Local RAG Workshop.
Generates realistic college workshop documents:
1. college_workshop_syllabus.pdf (Multi-page PDF via pymupdf)
2. ai_engineering_handbook.docx (Multi-section DOCX via python-docx)
"""

from pathlib import Path
import pymupdf  # PyMuPDF
import docx     # python-docx

CURRENT_DIR = Path(__file__).resolve().parent
CURRENT_DIR.mkdir(parents=True, exist_ok=True)


def generate_syllabus_pdf(output_path: Path):
    """Creates a 2-page college workshop syllabus PDF using PyMuPDF."""
    doc = pymupdf.open()

    # --- Page 1: Course Info, Overview, and Eligibility ---
    page1 = doc.new_page(width=595, height=842)  # A4 standard
    rect1 = pymupdf.Rect(50, 50, 545, 800)

    text_page1 = """COLLEGE OF ENGINEERING & ADVANCED TECHNOLOGY
DEPARTMENT OF COMPUTER SCIENCE & ARTIFICIAL INTELLIGENCE
ACADEMIC YEAR 2026-2027

COURSE SYLLABUS: CS-482 APPLIED RETRIEVAL-AUGMENTED GENERATION (RAG)
Workshop Title: Building Grounded Local AI Applications with Ollama & Sentence Transformers

1. COURSE OVERVIEW & OBJECTIVE
This intensive technical workshop provides undergraduate and graduate students with hands-on
experience in modern Retrieval-Augmented Generation (RAG) architectures. Students learn how to
ingest unstructured documents, create semantic embeddings locally, calculate cosine similarity,
and interface with open-weight language models running completely offline via Ollama.

2. ELIGIBILITY REQUIREMENTS & PREREQUISITES
To be eligible for workshop certification and credit, students must satisfy the following:
- Academic Standing: Enrolled in 3rd or 4th year B.Tech / M.Tech in Computer Science, Data Science,
  or allied engineering disciplines.
- Programming Prerequisites: Proficiency in Python 3.10+ (functions, data structures, list comprehensions).
- Mathematical Prerequisites: Fundamental understanding of Linear Algebra (vector dot products,
  Euclidean norm, and cosine angle) and basic Probability.
- Hardware Requirement: A student laptop with at least 8 GB RAM and 10 GB free disk space.
  A dedicated GPU is NOT required; all workshop embeddings and quantized models run on CPU.

3. WORKSHOP MODULE BREAKDOWN
Module 1: Unstructured Document Ingestion
- Extracting text from PDF and DOCX files.
- Document cleaning, ligature normalization, and handling non-printable characters.

Module 2: Semantic Chunking & Vector Embeddings
- Sliding-window chunking strategies and character overlap.
- Dense vector representations using the all-MiniLM-L6-v2 Sentence Transformer model (384 dimensions).
- L2-normalization of vector embeddings.
"""

    page1.insert_textbox(rect1, text_page1, fontsize=10, fontname="helv")

    # --- Page 2: Similarity, Retrieval, Grading, and Policies ---
    page2 = doc.new_page(width=595, height=842)
    rect2 = pymupdf.Rect(50, 50, 545, 800)

    text_page2 = """CS-482 WORKSHOP SYLLABUS — PAGE 2

Module 3: Vector Similarity & Top-K Retrieval
- Calculating directional cosine similarity between query vectors and chunk embeddings.
- Implementing Top-K ranking algorithms and minimum similarity threshold filtering.
- Mitigating hallucinations by rejecting ungrounded or out-of-domain queries.

Module 4: Local LLM Generation with Ollama
- Quantized local model inference using Ollama daemon.
- Primary Model: Qwen 2.5 (1.5B or 3B parameters) for concise, grounded answers.
- Fallback Model: Llama 3.2 (1B or 3B parameters) when Qwen is unavailable.
- Constructing injection-resistant grounded prompt templates with explicit citation rules.

4. GRADING & EVALUATION POLICY
Grading is strictly criterion-referenced. The total workshop score is 100 points:
- Active Workshop Lab Participation: 20% (hands-on pipeline execution).
- Mid-Workshop Vector Math Quiz: 30% (cosine similarity calculation and chunking theory).
- Capstone RAG Application Project: 50% (deploying a functional Chat-with-Docs UI).
Passing Grade: Students must achieve an aggregate score of at least 65% to earn the
Certificate of Completion.

5. ATTENDANCE & LATE SUBMISSION POLICIES
- Attendance Requirement: A minimum attendance of 80% across all workshop sessions is mandatory.
  Absences due to medical reasons require official department documentation within 48 hours.
- Late Submission Penalty: Capstone projects submitted past the deadline incur a penalty of
  10% deduction per day, up to a maximum of 3 days. Submissions older than 72 hours receive zero.

6. CONTACT & INSTRUCTOR DETAILS
- Workshop Lead: Dr. A. Vance, Senior AI Systems Architect & Adjunct Professor.
- Lab Teaching Assistant: T. K. Sharma (lab-ta@college.edu).
- Workshop Office Hours: Tuesdays & Thursdays, 4:00 PM - 5:30 PM, Lab Building 4, Room 302.
"""

    page2.insert_textbox(rect2, text_page2, fontsize=10, fontname="helv")

    doc.save(str(output_path))
    doc.close()
    print(f"Generated sample PDF: {output_path}")


def generate_handbook_docx(output_path: Path):
    """Creates an AI Engineering Technical Handbook DOCX using python-docx."""
    doc = docx.Document()

    # Title
    doc.add_heading("AI Engineering Technical Handbook: Local RAG Best Practices", level=0)
    doc.add_paragraph("Author: College Technical Workshop Series | Target: Undergraduate AI Engineers")

    # Section 1
    doc.add_heading("1. Understanding Retrieval-Augmented Generation", level=1)
    doc.add_paragraph(
        "Retrieval-Augmented Generation (RAG) is an architectural pattern that combines external knowledge "
        "retrieval with generative language models. Large language models possess broad reasoning capabilities "
        "but lack knowledge of private, recent, or proprietary enterprise documents. Rather than fine-tuning "
        "the model parameters—which is computationally expensive and quickly becomes obsolete—RAG dynamically "
        "retrieves relevant snippets from an indexed knowledge base and supplies them as reference context inside "
        "the prompt."
    )

    # Section 2
    doc.add_heading("2. Chunking Guidelines & Strategies", level=1)
    doc.add_paragraph(
        "Choosing chunk size involves balancing granularity and contextual integrity. A chunk size that is too small "
        "(e.g., 50 characters) loses semantic context, leading to poor embedding fidelity. Conversely, a chunk size "
        "that is too large (e.g., 5,000 characters) dilutes specific answers and exceeds model context budgets. "
        "For student laptops, a chunk size of 400 to 600 characters with an overlap of 50 characters provides an optimal "
        "trade-off between retrieval precision and memory footprint."
    )

    # Section 3
    doc.add_heading("3. Local Model Selection: Qwen vs Llama Policy", level=1)
    doc.add_paragraph(
        "In our workshop environment, all inference is conducted locally using Ollama without internet API dependencies. "
        "The model selection hierarchy prioritizes Qwen 2.5 (1.5B or 3B) due to its exceptional reasoning efficiency "
        "and strict adherence to prompt constraints. If Qwen is not detected in the local Ollama catalog, the pipeline "
        "automatically cascades to an available Llama 3.2 (1B or 3B) instance. If neither model is installed, the system "
        "halts gracefully with copy-paste installation commands rather than crashing."
    )

    # Section 4
    doc.add_heading("4. Preventing Hallucinations & Guarding Against Prompt Injection", level=1)
    doc.add_paragraph(
        "Hallucination control is enforced via a two-layer defense mechanism. First, a cosine similarity threshold gate "
        "(typically 0.30) blocks queries that have no semantic connection to the uploaded document. Second, the system prompt "
        "strictly confines the LLM to the provided context and instructs it to declare 'Information not found' when evidence "
        "is lacking. Furthermore, retrieved content is isolated within XML delimiters (<retrieved_context>) to prevent untrusted "
        "document text from overriding system-level instructions."
    )

    doc.save(str(output_path))
    print(f"Generated sample DOCX: {output_path}")


if __name__ == "__main__":
    pdf_path = CURRENT_DIR / "college_workshop_syllabus.pdf"
    docx_path = CURRENT_DIR / "ai_engineering_handbook.docx"

    generate_syllabus_pdf(pdf_path)
    generate_handbook_docx(docx_path)
